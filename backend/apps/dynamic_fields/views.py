from django.db.models import Count, Case, When, Value, IntegerField
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from .models import EntityField, EntityType, EntityRecord
from .serializers import EntityFieldSerializer, EntityTypeSerializer, EntityRecordSerializer
from .mixins import invalidate_fields_cache
from apps.workers.models import Worker
from apps.shifts.models import Shift


class EntityTypeViewSet(viewsets.ModelViewSet):
    serializer_class = EntityTypeSerializer
    permission_classes = [AllowAny]
    lookup_field = 'slug'

    def get_queryset(self):
        worker_count = Worker.objects.count()
        shift_count = Shift.objects.count()
        return EntityType.objects.annotate(
            _eav_count=Count('records'),
            record_count=Case(
                When(slug='worker', then=Value(worker_count)),
                When(slug='shift', then=Value(shift_count)),
                default='_eav_count',
                output_field=IntegerField(),
            ),
        ).order_by('order', 'name')

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if instance.is_system:
            return Response(
                {'error': 'No se pueden eliminar entidades del sistema'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        return super().destroy(request, *args, **kwargs)


class EntityRecordViewSet(viewsets.ModelViewSet):
    queryset = EntityRecord.objects.select_related('entity_type').all()
    serializer_class = EntityRecordSerializer
    permission_classes = [AllowAny]

    def get_queryset(self):
        qs = EntityRecord.objects.select_related('entity_type').all()
        entity_type = self.request.query_params.get('entity_type')
        if entity_type:
            qs = qs.filter(entity_type__slug=entity_type)
        return qs

    def destroy(self, request, *args, **kwargs):
        """Block deletion if other EAV records or workers reference this record."""
        instance = self.get_object()
        record_id = instance.id
        entity_slug = instance.entity_type_id  # slug of the entity being deleted

        dependents = self._find_dependents(entity_slug, record_id)
        if dependents:
            return Response(
                {
                    'error': 'No se puede eliminar: hay registros que dependen de este.',
                    'dependents': dependents,
                },
                status=status.HTTP_409_CONFLICT,
            )
        return super().destroy(request, *args, **kwargs)

    def _find_dependents(self, entity_slug, record_id):
        """
        Search all EntityField definitions that point to `entity_slug` via
        entity_select or multi_entity_select. Then check if any EntityRecord
        or Worker references the given record_id in their data/custom_data.
        Returns a list of {entity: <slug>, label: <display>, count: N} dicts.
        """
        # Find all fields (across any entity) that target this entity type
        referencing_fields = EntityField.objects.filter(
            target_entity=entity_slug,
            field_type__in=['entity_select', 'multi_entity_select'],
            active=True,
        )

        dependents = []
        str_id = str(record_id)

        for field in referencing_fields:
            owner_entity = field.entity_type  # slug of the entity that HAS this field

            if owner_entity == 'worker':
                # Workers use custom_data (JSON) on the Worker model
                count = 0
                for w in Worker.objects.filter(active=True):
                    val = (w.custom_data or {}).get(field.key)
                    if self._value_references(val, str_id):
                        count += 1
                if count > 0:
                    et = EntityType.objects.filter(slug='worker').first()
                    label = et.name if et else 'Empleados'
                    dependents.append({
                        'entity': 'worker',
                        'field': field.key,
                        'label': f'{count} {label} (campo "{field.label}")',
                        'count': count,
                    })
            else:
                # Generic EAV records
                count = 0
                for rec in EntityRecord.objects.filter(entity_type=owner_entity):
                    val = (rec.data or {}).get(field.key)
                    if self._value_references(val, str_id):
                        count += 1
                if count > 0:
                    et = EntityType.objects.filter(slug=owner_entity).first()
                    label = et.name if et else owner_entity
                    dependents.append({
                        'entity': owner_entity,
                        'field': field.key,
                        'label': f'{count} {label} (campo "{field.label}")',
                        'count': count,
                    })

        return dependents

    @staticmethod
    def _value_references(val, str_id):
        """Check if a custom_data value (scalar, list, list-of-dicts) references the given ID."""
        if val is None:
            return False
        if isinstance(val, list):
            for item in val:
                if isinstance(item, dict):
                    for k in ('entity_id', 'id', 'value'):
                        if k in item and str(item[k]) == str_id:
                            return True
                elif str(item) == str_id:
                    return True
            return False
        return str(val) == str_id

    def list(self, request, *args, **kwargs):
        """Override list to inject model-backed records for shift/worker."""
        entity_type = request.query_params.get('entity_type')
        if entity_type == 'shift':
            return Response(self._shift_records())
        if entity_type == 'worker':
            return Response(self._worker_records())
        return super().list(request, *args, **kwargs)

    @action(detail=False, methods=['get'], url_path='export')
    def export_excel(self, request):
        """Export entity records to .xlsx. Query param: entity_type=<slug>"""
        import io
        from django.http import HttpResponse
        from datetime import date as _date

        entity_type_slug = request.query_params.get('entity_type', '')
        if not entity_type_slug:
            return Response({'error': 'entity_type param required'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({'error': 'openpyxl no instalado'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        from apps.catalog.models import KindValue
        from apps.shifts.models import Shift as ShiftModel

        # ── Load entity type + fields ────────────────────────────────────
        try:
            et = EntityType.objects.get(slug=entity_type_slug)
        except EntityType.DoesNotExist:
            return Response({'error': 'entity_type not found'}, status=status.HTTP_404_NOT_FOUND)

        worker_fields = list(EntityField.objects.filter(entity_type=entity_type_slug, active=True).order_by('order'))
        display_key = et.display_field or 'nombre'

        # ── Build label lookups for entity_select fields ─────────────────
        entity_label = {}   # field.key → {record_id: label}
        catalog_label = {}  # field.key → {code: label}

        for field in worker_fields:
            if field.field_type in ('entity_select', 'multi_entity_select') and field.target_entity:
                dk = field.display_key or 'nombre'
                lmap = {}
                if field.target_entity == 'shift':
                    for s in ShiftModel.objects.all():
                        lmap[s.id] = s.name
                elif field.target_entity == 'worker':
                    for w in Worker.objects.all():
                        lmap[w.id] = w.name
                else:
                    for r in EntityRecord.objects.filter(entity_type=field.target_entity):
                        lmap[r.id] = str(r.data.get(dk) or r.data.get('nombre') or r.data.get('name') or r.id)
                entity_label[field.key] = lmap
            elif field.field_type in ('catalog_select', 'multi_catalog_select') and field.kind_code:
                catalog_label[field.key] = {
                    kv.code: kv.label
                    for kv in KindValue.objects.filter(kind__code=field.kind_code, active=True)
                }

        def _resolve(field, raw):
            if raw is None or raw == '':
                return ''
            if field.field_type == 'boolean':
                return 'Sí' if raw is True or str(raw).lower() in ('true', '1') else 'No'
            if field.field_type in ('entity_select', 'multi_entity_select'):
                lmap = entity_label.get(field.key, {})
                ids = raw if isinstance(raw, list) else [raw]
                labels = []
                for item in ids:
                    rid = item.get('value', item) if isinstance(item, dict) else item
                    try: rid = int(rid)
                    except (ValueError, TypeError): rid = None
                    labels.append(lmap.get(rid, str(rid) if rid else ''))
                return ', '.join(l for l in labels if l)
            if field.field_type in ('catalog_select', 'multi_catalog_select'):
                clmap = catalog_label.get(field.key, {})
                codes = raw if isinstance(raw, list) else [raw]
                return ', '.join(clmap.get(str(c), str(c)) for c in codes if c)
            if isinstance(raw, list):
                return ', '.join(str(x) for x in raw)
            return str(raw)

        # ── Load records ─────────────────────────────────────────────────
        if entity_type_slug == 'shift':
            records_data = [
                {
                    'id': s.id,
                    **(s.custom_data or {}),
                    # Model fields always override stale custom_data copies
                    'nombre': s.name,
                    'hora_llegada': str(s.start_time)[:5] if s.start_time else '',
                    'hora_salida': str(s.end_time)[:5] if s.end_time else '',
                }
                for s in ShiftModel.objects.order_by('name')
            ]
        elif entity_type_slug == 'worker':
            records_data = [
                {'id': w.id, 'nombre': w.name, **(w.custom_data or {})}
                for w in Worker.objects.order_by('name')
            ]
        else:
            records_data = [
                {'id': r.id, **r.data}
                for r in EntityRecord.objects.filter(entity_type=entity_type_slug).order_by('id')
            ]

        # ── Skip display_field from schema columns (already in col 2) ────
        # Avoids duplicating the "Nombre" column when it's also an EntityField.
        schema_fields = [f for f in worker_fields if f.key != display_key]

        # ── Build workbook ───────────────────────────────────────────────
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = et.name[:31]

        hdr_fill = PatternFill('solid', fgColor='1F497D')
        hdr_font = Font(bold=True, color='FFFFFF')

        def _hdr(col, text, width=20):
            c = ws.cell(row=1, column=col, value=text)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal='center')
            ws.column_dimensions[get_column_letter(col)].width = max(width, len(text) + 4)

        _hdr(1, 'ID', 8)
        _hdr(2, 'Nombre', 30)
        for i, field in enumerate(schema_fields, start=3):
            _hdr(i, field.label)

        ws.freeze_panes = 'A2'

        for row_idx, rec in enumerate(records_data, start=2):
            ws.cell(row=row_idx, column=1, value=rec.get('id'))
            ws.cell(row=row_idx, column=2, value=str(rec.get(display_key) or rec.get('nombre') or rec.get('name') or ''))
            for col_idx, field in enumerate(schema_fields, start=3):
                ws.cell(row=row_idx, column=col_idx, value=_resolve(field, rec.get(field.key)))

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f'{entity_type_slug}-{_date.today().isoformat()}.xlsx'
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename={filename}'
        return response


    def _shift_records(self):
        """Turnos para los desplegables `entity_select` (turno base/alternativo/
        sábado…).

        Excluye los turnos extra ya caducados (`valid_to` < hoy), misma regla
        que `?active=true` en ShiftViewSet: no se ofrecen para NUEVAS
        asignaciones, pero siguen existiendo como registro histórico y se ven
        en la vista de admin de turnos. Los turnos normales (sin vigencia) y los
        que aún no han empezado pasan siempre.
        """
        from datetime import date as _date

        from django.db.models import Q

        from apps.shifts.models import Shift
        qs = Shift.objects.filter(
            Q(valid_to__isnull=True) | Q(valid_to__gte=_date.today())
        ).order_by('name')

        records = []
        for s in qs:
            records.append({
                'id': s.id,
                'entity_type': 'shift',
                'entity_type_slug': 'shift',
                'data': {
                    **(s.custom_data or {}),
                    'nombre': s.name,
                    'hora_llegada': s.start_time.strftime('%H:%M') if s.start_time else '',
                    'hora_salida': s.end_time.strftime('%H:%M') if s.end_time else '',
                },
                'field_schema': [],
                'created_at': '',
                'updated_at': '',
            })
        return records

    def _worker_records(self):
        from apps.workers.models import Worker
        from apps.workers.serializers import WorkerSerializer
        workers = Worker.objects.order_by('id')
        serialized = WorkerSerializer(workers, many=True).data
        records = []
        for w in serialized:
            records.append({
                'id': w['id'],
                'entity_type': 'worker',
                'entity_type_slug': 'worker',
                'data': {
                    **(w.get('custom_data') or {}),
                    'nombre': w['name'],
                    'activo': w.get('active', True),
                },
                'field_schema': w.get('field_schema', []),
                'created_at': '',
                'updated_at': '',
            })
        return records


class EntityFieldViewSet(viewsets.ModelViewSet):
    queryset = EntityField.objects.all()
    serializer_class = EntityFieldSerializer
    permission_classes = [AllowAny]

    def get_queryset(self):
        qs = EntityField.objects.all()
        entity_type = self.request.query_params.get('entity_type')
        if entity_type:
            qs = qs.filter(entity_type=entity_type)
        # Los campos desactivados no se ofrecen a los formularios (?active=1).
        # Es la forma de retirar un campo sin borrarlo: el campo y los datos ya
        # guardados siguen en BD, simplemente deja de pintarse. Sin el parámetro
        # se devuelven todos, para que el panel de administración pueda
        # gestionarlos (incluido reactivarlos).
        if self.request.query_params.get('active') in ('1', 'true', 'True'):
            qs = qs.filter(active=True)
        return qs

    def perform_create(self, serializer):
        instance = serializer.save()
        invalidate_fields_cache(instance.entity_type)

    def perform_update(self, serializer):
        instance = serializer.save()
        invalidate_fields_cache(instance.entity_type)

    def perform_destroy(self, instance):
        entity_type = instance.entity_type
        instance.delete()
        invalidate_fields_cache(entity_type)

    @action(detail=False, methods=['get'], url_path='by-entity/(?P<entity_type>[a-z_]+)')
    def by_entity(self, request, entity_type=None):
        fields = EntityField.objects.filter(entity_type=entity_type, active=True).order_by('order')
        serializer = self.get_serializer(fields, many=True)
        return Response(serializer.data)
