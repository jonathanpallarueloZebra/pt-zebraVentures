import logging
from datetime import date, datetime, timedelta

from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from .models import AbsenceType, AbsenceRequest
from .serializers import AbsenceTypeSerializer, AbsenceRequestSerializer
from apps.rest_days.models import RestDay
from apps.email_service.service import EmailService

logger = logging.getLogger(__name__)
User = get_user_model()
email_service = EmailService()


class AbsenceTypeViewSet(viewsets.ModelViewSet):
    queryset = AbsenceType.objects.all()
    serializer_class = AbsenceTypeSerializer
    permission_classes = [AllowAny]

    def get_queryset(self):
        if self.action == 'list':
            return AbsenceType.objects.filter(active=True)
        return AbsenceType.objects.all()


class AbsenceRequestViewSet(viewsets.ModelViewSet):
    queryset = AbsenceRequest.objects.select_related('worker', 'type', 'reviewed_by').all()
    serializer_class = AbsenceRequestSerializer
    permission_classes = [AllowAny]

    # ── Excel: plantilla, exportacion e importacion ──────────────
    #
    # Mismo patron que WorkerViewSet (import-template / export / import-excel):
    # openpyxl, cabeceras azules, desplegables por validacion de datos y un
    # resultado {created, errors, total_rows} para que el front pinte el aviso.

    @staticmethod
    def _codigo_de(worker):
        """Codigo del empleado (custom_data['codigo']). Es el identificador que
        usa la plantilla: los 260 empleados vigentes lo tienen y no hay
        duplicados, asi que sirve para localizarlos sin ambigüedad."""
        return str((worker.custom_data or {}).get('codigo') or '').strip()

    @staticmethod
    def _tipos_disponibles():
        """Etiquetas del catalogo 'absence_type' (AdminZebra > Catalogo).

        Se leen del catalogo y no de AbsenceType para que la plantilla ofrezca
        los tipos que el cliente ve en la aplicacion, incluidos los que anada
        el mismo. Si el catalogo no existiera, se cae a AbsenceType.
        """
        from apps.catalog.models import KindValue
        etiquetas = list(
            KindValue.objects.filter(kind__code='absence_type', active=True)
            .order_by('order').values_list('label', flat=True)
        )
        if etiquetas:
            return etiquetas
        return list(AbsenceType.objects.filter(active=True).values_list('name', flat=True))

    @action(detail=False, methods=['get'], url_path='import-template')
    def import_template(self, request):
        """Excel de ejemplo con los empleados vigentes ya rellenados (AC-3/AC-4).

        Columnas: Codigo | Nombre | Tipo | Fecha inicio | Fecha fin |
        Indefinida | Observaciones.

        Codigo y Nombre vienen precargados con los empleados ACTIVOS en el
        momento de la descarga, para que solo haya que rellenar la ausencia en
        la fila de quien la tenga. Las filas que se dejen sin datos se ignoran
        al importar.
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.datavalidation import DataValidation
        except ImportError:
            return Response(
                {'error': 'openpyxl no esta instalado en el servidor.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )
        from apps.workers.models import Worker

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ausencias'

        hdr_fill = PatternFill('solid', fgColor='1F497D')
        hdr_font = Font(bold=True, color='FFFFFF')
        gris = Font(color='808080')

        columnas = [
            ('Codigo', 14), ('Nombre', 34), ('Tipo de ausencia *', 24),
            ('Fecha inicio *', 16), ('Fecha fin', 16),
            ('Indefinida (SI/NO)', 18), ('Observaciones', 40),
        ]
        for i, (texto, ancho) in enumerate(columnas, start=1):
            c = ws.cell(row=1, column=i, value=texto)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal='center')
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = 'A2'

        # Empleados vigentes en el momento de la descarga.
        empleados = [
            w for w in Worker.objects.filter(active=True).order_by('name')
            if self._codigo_de(w)
        ]
        for fila, w in enumerate(empleados, start=2):
            ws.cell(row=fila, column=1, value=self._codigo_de(w)).font = gris
            ws.cell(row=fila, column=2, value=w.name).font = gris

        ultima = len(empleados) + 1

        # Desplegable de tipos, con los valores reales del catalogo.
        tipos = self._tipos_disponibles()
        if tipos:
            hoja_ref = wb.create_sheet('Referencia')
            for i, t in enumerate(tipos, start=1):
                hoja_ref.cell(row=i, column=1, value=t)
            hoja_ref.sheet_state = 'hidden'
            dv = DataValidation(
                type='list',
                formula1=f'=Referencia!$A$1:$A${len(tipos)}',
                allow_blank=True,
                showDropDown=False,
            )
            dv.error = 'Elige un tipo de la lista.'
            ws.add_data_validation(dv)
            dv.add(f'C2:C{max(ultima, 2)}')

        # SI/NO para la indefinida
        dv_bool = DataValidation(
            type='list', formula1='"SI,NO"', allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv_bool)
        dv_bool.add(f'F2:F{max(ultima, 2)}')

        # Instrucciones, en una hoja aparte para no estorbar a la lectura.
        guia = wb.create_sheet('Instrucciones')
        for i, linea in enumerate([
            'Como rellenar esta plantilla',
            '',
            '1. Codigo y Nombre vienen ya rellenos con los empleados activos.',
            '   No hace falta tocarlos: localiza la fila del empleado y rellena',
            '   su ausencia a la derecha.',
            '2. Solo se importan las filas que tengan datos de ausencia.',
            '   Las filas que dejes sin rellenar se ignoran.',
            '3. Tipo de ausencia: elige uno de la lista del desplegable.',
            '4. Fechas en formato DD/MM/AAAA (o celda con formato fecha).',
            '5. Fecha fin es obligatoria salvo que Indefinida sea SI.',
            '6. Si una fila esta mal, se avisa de esa fila y el resto se importa.',
            '',
            'Cada fila crea una ausencia NUEVA. Importar dos veces el mismo',
            'fichero crea las ausencias dos veces.',
        ], start=1):
            celda = guia.cell(row=i, column=1, value=linea)
            if i == 1:
                celda.font = Font(bold=True, size=13)
        guia.column_dimensions['A'].width = 78

        import io as _io
        buff = _io.BytesIO()
        wb.save(buff)
        buff.seek(0)
        resp = HttpResponse(
            buff.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = 'attachment; filename="plantilla_ausencias.xlsx"'
        return resp

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Exporta el listado de ausencias a Excel (AC-1)."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response(
                {'error': 'openpyxl no esta instalado en el servidor.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ausencias'
        hdr_fill = PatternFill('solid', fgColor='1F497D')
        hdr_font = Font(bold=True, color='FFFFFF')

        columnas = [
            ('Codigo', 14), ('Nombre', 34), ('Tipo de ausencia', 24),
            ('Fecha inicio', 16), ('Fecha fin', 16), ('Indefinida', 12),
            ('Observaciones', 40),
        ]
        for i, (texto, ancho) in enumerate(columnas, start=1):
            c = ws.cell(row=1, column=i, value=texto)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal='center')
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = 'A2'

        # Se exporta lo mismo que ve la pantalla (respeta los filtros de la URL).
        for fila, a in enumerate(self.filter_queryset(self.get_queryset()), start=2):
            ws.cell(row=fila, column=1, value=self._codigo_de(a.worker))
            ws.cell(row=fila, column=2, value=a.worker.name)
            ws.cell(row=fila, column=3, value=a.type.name if a.type_id else '')
            ws.cell(row=fila, column=4, value=a.start_date)
            # Una indefinida no tiene fin: la celda se deja vacia (no un texto)
            # para que el fichero exportado se pueda volver a importar tal cual.
            ws.cell(row=fila, column=5, value=a.end_date if a.end_date else '')
            ws.cell(row=fila, column=6, value='SI' if a.indefinite else 'NO')
            ws.cell(row=fila, column=7, value=a.reason or '')
            for col in (4, 5):
                ws.cell(row=fila, column=col).number_format = 'DD/MM/YYYY'

        import io as _io
        buff = _io.BytesIO()
        wb.save(buff)
        buff.seek(0)
        resp = HttpResponse(
            buff.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        hoy = timezone.now().date().isoformat()
        resp['Content-Disposition'] = f'attachment; filename="ausencias-{hoy}.xlsx"'
        return resp

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        """Importa ausencias desde el Excel de la plantilla (AC-2/AC-5/AC-6).

        Cada fila con datos crea una ausencia NUEVA: no actualiza las que ya
        existen, porque una ausencia no tiene clave natural (un empleado puede
        tener varias del mismo tipo). Si se importa dos veces el mismo fichero
        se crean dos veces, asi que se avisa de los solapamientos.

        Las filas sin datos de ausencia se ignoran: la plantilla trae los 260
        empleados precargados y lo normal es rellenar solo unas pocas.

        Con ?revisar=1 (o revisar=true en el formulario) NO escribe nada: solo
        valida y devuelve que pasaria. Lo usa la pantalla al soltar el fichero,
        para avisar de los fallos ANTES de importar.
        """
        try:
            import openpyxl
        except ImportError:
            return Response(
                {'error': 'openpyxl no esta instalado en el servidor.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )
        from apps.workers.models import Worker
        from .serializers import resolver_absence_type

        solo_revisar = str(
            request.query_params.get('revisar')
            or request.data.get('revisar')
            or ''
        ).lower() in ('1', 'true', 'si', 'yes')

        fichero = request.FILES.get('file')
        if not fichero:
            return Response({'error': 'No se ha recibido ningun fichero.'},
                            status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(fichero, data_only=True)
        except Exception as exc:
            return Response({'error': f'No se pudo leer el Excel: {exc}'},
                            status=status.HTTP_400_BAD_REQUEST)

        ws = wb['Ausencias'] if 'Ausencias' in wb.sheetnames else wb.active

        # Empleados por codigo. Se aceptan tambien los inactivos: una ausencia
        # historica puede ser de alguien que ya no esta de alta.
        por_codigo = {}
        for w in Worker.objects.all():
            cod = self._codigo_de(w)
            if cod:
                por_codigo[cod.lower()] = w

        def _texto(v):
            if v is None:
                return ''
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v).strip()

        def _fecha(v, etiqueta, fallos):
            """Acepta celda con formato fecha o texto DD/MM/AAAA (y AAAA-MM-DD)."""
            if v is None or _texto(v) == '':
                return None
            if isinstance(v, datetime):
                return v.date()
            if isinstance(v, date):
                return v
            txt = _texto(v)
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
                try:
                    return datetime.strptime(txt, fmt).date()
                except ValueError:
                    continue
            fallos.append(f'{etiqueta} "{txt}" no es una fecha valida (usa DD/MM/AAAA)')
            return None

        def _si_no(v):
            return _texto(v).lower() in ('si', 'sí', 'true', '1', 'x', 'yes')

        creadas = 0
        errores = []
        filas_con_datos = 0

        for i, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            def col(idx, _fila=fila):
                return _fila[idx] if len(_fila) > idx else None

            codigo = _texto(col(0))
            nombre = _texto(col(1))
            tipo = _texto(col(2))
            f_ini_raw, f_fin_raw = col(3), col(4)
            indefinida = _si_no(col(5))
            observaciones = _texto(col(6))

            # AC-5: una fila solo cuenta si trae DATOS DE AUSENCIA. Codigo y
            # nombre vienen precargados en las 260 filas de la plantilla, asi
            # que mirarlos aqui haria que se procesaran todas.
            hay_datos = any([
                tipo, _texto(f_ini_raw), _texto(f_fin_raw),
                indefinida, observaciones,
            ])
            if not hay_datos:
                continue
            filas_con_datos += 1

            etiqueta = f'Fila {i}' + (f' ({nombre})' if nombre else '')
            fallos = []

            trabajador = por_codigo.get(codigo.lower()) if codigo else None
            if not trabajador:
                fallos.append(
                    f'no se encuentra el empleado con codigo "{codigo}"'
                    if codigo else 'falta el codigo del empleado')

            if not tipo:
                fallos.append('falta el tipo de ausencia')

            inicio = _fecha(f_ini_raw, 'Fecha inicio', fallos)
            fin = _fecha(f_fin_raw, 'Fecha fin', fallos)
            if inicio is None and not any('Fecha inicio' in f for f in fallos):
                fallos.append('falta la fecha de inicio')
            if not indefinida and fin is None and not any('Fecha fin' in f for f in fallos):
                fallos.append('falta la fecha de fin (o marca Indefinida = SI)')
            if inicio and fin and fin < inicio:
                fallos.append('la fecha de fin es anterior a la de inicio')

            # AC-6: se acumula el fallo de esta fila y se sigue con las demas.
            if fallos:
                errores.append(f'{etiqueta}: ' + '; '.join(fallos))
                continue

            # En modo revision la fila es valida y no se escribe nada.
            if solo_revisar:
                creadas += 1
                if AbsenceRequest.objects.filter(
                        worker=trabajador, start_date=inicio).exists():
                    errores.append(
                        f'{etiqueta}: ya hay otra ausencia de este empleado que '
                        f'empieza el mismo dia (se creara duplicada)')
                continue

            try:
                objeto_tipo = resolver_absence_type(tipo)
                solapa = AbsenceRequest.objects.filter(
                    worker=trabajador, start_date=inicio).exists()
                AbsenceRequest.objects.create(
                    worker=trabajador,
                    type=objeto_tipo,
                    start_date=inicio,
                    end_date=None if indefinida else fin,
                    indefinite=indefinida,
                    reason=observaciones,
                    # La empresa registra directo, sin moderacion: nace aprobada
                    # para que el planificador la tenga en cuenta.
                    status='approved',
                    custom_data={'tipo': objeto_tipo.name if objeto_tipo else tipo},
                )
                creadas += 1
                if solapa:
                    errores.append(
                        f'{etiqueta}: creada, pero ya habia otra ausencia de este '
                        f'empleado que empieza el mismo dia (revisa si esta duplicada)')
            except Exception as exc:
                errores.append(f'{etiqueta}: no se pudo crear - {exc}')

        return Response({
            # En modo revision, 'created' es cuantas se crearian.
            'created': creadas,
            'errors': errores,
            'total_rows': filas_con_datos,
            'dry_run': solo_revisar,
        })

    def get_queryset(self):
        qs = super().get_queryset()
        status_filter = self.request.query_params.get('status')
        worker_filter = self.request.query_params.get('worker')
        if status_filter:
            qs = qs.filter(status=status_filter)
        if worker_filter:
            qs = qs.filter(worker_id=worker_filter)
        return qs

    def perform_create(self, serializer):
        instance = serializer.save()
        # Notify admin/HR staff about the new request
        self._notify_admins_new_request(instance)

    def perform_update(self, serializer):
        """Edición de una ausencia (AC-5). Si se le añade una fecha de fin real
        a una ausencia indefinida, deja de ser indefinida y, si ya estaba
        aprobada, se recalculan los RestDay para que el empleado vuelva a estar
        disponible a partir del día siguiente al nuevo fin.
        """
        was_open_ended = serializer.instance.is_open_ended
        instance = serializer.save()

        if was_open_ended and instance.end_date is not None:
            # El serializer ya fuerza indefinite→end_date=None; aquí el caller
            # ha mandado indefinite=False + una fecha de fin, cerrando la baja.
            if instance.status == 'approved':
                self._sync_rest_days(instance)

    def _sync_rest_days(self, absence):
        """Reconstruye los RestDay de la ausencia según su rango actual: crea
        los que falten dentro de [start, end] y elimina los posteriores al fin
        (que quedaron de cuando era indefinida)."""
        # Eliminar días de descanso de esta ausencia posteriores al nuevo fin.
        RestDay.objects.filter(
            worker=absence.worker,
            reason=absence.type.name,
            date__gt=absence.end_date,
        ).delete()
        # Asegurar los del rango vigente.
        current = absence.start_date
        while current <= absence.end_date:
            RestDay.objects.get_or_create(
                worker=absence.worker,
                date=current,
                defaults={'reason': absence.type.name},
            )
            current += timedelta(days=1)

    @action(detail=True, methods=['patch'])
    def approve(self, request, pk=None):
        absence = self.get_object()
        if absence.status != 'pending':
            return Response(
                {'error': 'Solo se pueden aprobar solicitudes pendientes.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        absence.status = 'approved'
        absence.reviewed_by = request.user if request.user.is_authenticated else None
        absence.reviewed_at = timezone.now()
        absence.save()

        # Create RestDay entries for each date in the range. Una ausencia
        # indefinida (sin fecha de fin) NO genera RestDay concretos —su bloqueo
        # lo aplica el generador desde start_date en adelante (AC-4)—; cuando se
        # cierre con una fecha de fin real, perform_update creará los RestDay.
        if absence.end_date is not None:
            current = absence.start_date
            while current <= absence.end_date:
                RestDay.objects.get_or_create(
                    worker=absence.worker,
                    date=current,
                    defaults={'reason': absence.type.name},
                )
                current += timedelta(days=1)

        # Notify worker
        self._notify_worker_decision(absence, 'aprobada')

        serializer = self.get_serializer(absence)
        return Response(serializer.data)

    @action(detail=True, methods=['patch'])
    def reject(self, request, pk=None):
        absence = self.get_object()
        if absence.status != 'pending':
            return Response(
                {'error': 'Solo se pueden rechazar solicitudes pendientes.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        absence.status = 'rejected'
        absence.reviewed_by = request.user if request.user.is_authenticated else None
        absence.reviewed_at = timezone.now()
        absence.save()

        # Notify worker
        self._notify_worker_decision(absence, 'rechazada')

        serializer = self.get_serializer(absence)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='pending-count')
    def pending_count(self, request):
        count = AbsenceRequest.objects.filter(status='pending').count()
        return Response({'count': count})

    def _notify_admins_new_request(self, absence):
        """Send email to all staff users about a new absence request."""
        admins = User.objects.filter(is_staff=True, is_active=True)
        worker_name = absence.worker.name if absence.worker else 'Desconocido'
        for admin in admins:
            email_service.send_plain(
                to_email=admin.email,
                subject=f'Nueva solicitud de ausencia - {worker_name}',
                body=(
                    f'{worker_name} ha solicitado una ausencia.\n'
                    f'Tipo: {absence.type.name}\n'
                    f'Desde: {absence.start_date}\n'
                    f'Hasta: {absence.end_date if absence.end_date else "Indefinida"}\n'
                    f'Motivo: {absence.reason or "-"}\n\n'
                    f'Revisa la solicitud en el panel de ausencias.'
                ),
            )
        logger.info('Notified %d admins about absence request %d', admins.count(), absence.pk)

    def _notify_worker_decision(self, absence, decision):
        """Send email to worker about their absence request decision."""
        if not absence.worker or not absence.worker.user:
            return
        worker_email = absence.worker.user.email
        email_service.send_plain(
            to_email=worker_email,
            subject=f'Solicitud de ausencia {decision}',
            body=(
                f'Hola {absence.worker.name},\n\n'
                f'Tu solicitud de ausencia ({absence.type.name}) '
                + (f'del {absence.start_date} al {absence.end_date} '
                   if absence.end_date else f'desde el {absence.start_date} (indefinida) ')
                + f'ha sido {decision}.\n'
            ),
        )
        logger.info('Notified worker %s about %s decision', worker_email, decision)
